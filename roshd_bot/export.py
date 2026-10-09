"""خروجی‌ها: اکسل آمار تکالیف و فایل zip پاسخ‌های یک تکلیف."""
from __future__ import annotations

import re
import zipfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from .rules import EXCUSED, LATE, MISSING, ON_TIME, OPEN, STREAK_LEN
from .timeutil import fa, fmt_dt, from_db, jalali_str

LOGO = Path(__file__).parent / "assets" / "logo.png"
FONT = "Tahoma"
GREEN = "8BC888"
FILL = {
    ON_TIME: "D9F2D9", LATE: "FFE7B3", MISSING: "F8C9C9", OPEN: "E8EEF7", EXCUSED: "E6DDF3",
}
LABEL = {ON_TIME: "به‌موقع", LATE: "با تاخیر", MISSING: "ارسال‌نشده", OPEN: "در مهلت", EXCUSED: "عذر موجه"}
KIND_FA = {"regular": "اصلی", "optional": "اختیاری", "bonus": "امتیازی"}
thin = Side(style="thin", color="BBBBBB")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def _style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4E9A5B")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER


def _finish(ws, widths):
    ws.sheet_view.rightToLeft = True
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    for row in ws.iter_rows():
        for cell in row:
            if cell.font.name != FONT:
                cell.font = Font(name=FONT, bold=cell.font.bold, color=cell.font.color)
            if cell.row > 4 and cell.alignment.horizontal is None:
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.border = BORDER


def _title(ws, text, ncols):
    ws.row_dimensions[1].height = 24
    if LOGO.exists():
        img = XLImage(str(LOGO))
        img.height, img.width = 60, 69
        ws.add_image(img, "A1")
    ws.cell(row=2, column=3, value=text).font = Font(name=FONT, bold=True, size=14, color="2F6B3A")
    ws.merge_cells(start_row=2, start_column=3, end_row=2, end_column=max(ncols, 4))
    for r in (1, 2, 3):
        ws.row_dimensions[r].height = 24


def build_excel(repo, cfg, path: Path, now) -> Path:
    students = repo.students()
    assigns = list(reversed(repo.q("SELECT * FROM assignments ORDER BY deadline DESC")))
    targeted = {a["id"]: {s["id"] for s in repo.students_for(a["id"])} for a in assigns}
    outcome = {}
    for a in assigns:
        for s in students:
            if s["id"] in targeted[a["id"]] and repo.counts_for(a, s):
                outcome[(a["id"], s["id"])] = repo.evaluate(a, s["id"], now)

    wb = Workbook()
    # --- خلاصه
    ws = wb.active
    ws.title = "خلاصه دانش‌آموزان"
    heads = ["ردیف", "دانش‌آموز", "گروه‌ها", "تکالیف اصلی", "به‌موقع", "با تاخیر", "ارسال‌نشده", "عذر موجه",
             "درصد به‌موقع", "امتیاز", "تعداد نشان", f"پیشرفت تا نشان بعدی (از {fa(STREAK_LEN)})"]
    _title(ws, f"آمار تکالیف — {cfg.org_name} — {fa(jalali_str(now.astimezone(cfg.tz)))}", len(heads))
    for c, h in enumerate(heads, 1):
        ws.cell(row=4, column=c, value=h)
    _style_header(ws, 4, len(heads))
    for i, s in enumerate(students, 1):
        sm = repo.summary(s["id"], now)
        reg = [(a, outcome[(a["id"], s["id"])][0]) for a in assigns
               if a["kind"] == "regular" and (a["id"], s["id"]) in outcome]
        exc = sum(1 for _, o in reg if o.status == EXCUSED)
        denom = sm.on_time + sm.late + sm.missing
        ws.append([i, s["full_name"], s["groups"] or "", len(reg), sm.on_time, sm.late, sm.missing, exc,
                   round(sm.on_time / denom, 3) if denom else None, sm.points, len(repo.student_badges(s["id"])), sm.run])
        ws.cell(row=ws.max_row, column=9).number_format = "0%"
    ws.freeze_panes = "C5"
    _finish(ws, [6, 22, 22, 12, 10, 10, 12, 10, 12, 9, 10, 18])

    # --- ماتریس
    ws = wb.create_sheet("ماتریس تکالیف")
    heads = ["دانش‌آموز"] + [f"{a['title']}\n{fa(jalali_str(from_db(a['deadline']).astimezone(cfg.tz)))}\n({KIND_FA[a['kind']]})" for a in assigns]
    _title(ws, "وضعیت هر دانش‌آموز در هر تکلیف", len(heads))
    for c, h in enumerate(heads, 1):
        ws.cell(row=4, column=c, value=h)
    _style_header(ws, 4, len(heads))
    ws.row_dimensions[4].height = 48
    for s in students:
        row = [s["full_name"]]
        for a in assigns:
            k = (a["id"], s["id"])
            if k not in outcome:
                row.append("—")
                continue
            o = outcome[k][0]
            t = LABEL[o.status]
            if o.status == LATE:
                t += f"\n{fa(o.late_seconds // 60)} دقیقه"
            row.append(t)
        ws.append(row)
        for c, a in enumerate(assigns, 2):
            k = (a["id"], s["id"])
            if k in outcome:
                ws.cell(row=ws.max_row, column=c).fill = PatternFill("solid", fgColor=FILL[outcome[k][0].status])
    ws.freeze_panes = "B5"
    _finish(ws, [22] + [16] * len(assigns))

    # --- جزئیات
    ws = wb.create_sheet("جزئیات ارسال")
    heads = ["تکلیف", "نوع", "مهلت", "دانش‌آموز", "وضعیت", "زمان ارسال", "تاخیر (دقیقه)", "تعداد ارسال", "بازخورد مدیر"]
    _title(ws, "جزئیات ارسال‌ها", len(heads))
    for c, h in enumerate(heads, 1):
        ws.cell(row=4, column=c, value=h)
    _style_header(ws, 4, len(heads))
    for a in assigns:
        for s in students:
            k = (a["id"], s["id"])
            if k not in outcome:
                continue
            o, sub_at, dl = outcome[k]
            finals = repo.final_attempts(a["id"], s["id"])
            fb = next((f["feedback"] for f in reversed(finals) if f["feedback"]), "")
            label = LABEL[o.status] + (" (مهلت ارفاقی)" if o.in_grace else "")
            ws.append([a["title"], KIND_FA[a["kind"]], fmt_dt(dl, cfg.tz), s["full_name"], label,
                       fmt_dt(sub_at, cfg.tz) if sub_at else "",
                       round(o.late_seconds / 60) if o.status == LATE else None, len(finals), fb])
            ws.cell(row=ws.max_row, column=5).fill = PatternFill("solid", fgColor=FILL[o.status])
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:{ws.cell(row=4, column=len(heads)).column_letter}{max(ws.max_row, 5)}"
    _finish(ws, [24, 10, 26, 22, 20, 26, 12, 10, 36])

    wb.save(path)
    return path


def _safe(name: str) -> str:
    return re.sub(r'[\\/:*?"<>|]+', "_", name).strip() or "file"


async def build_zip(repo, bale, cfg, aid: int, path: Path):
    """آخرین ارسال هر دانش‌آموز برای یک تکلیف. خروجی: (مسیر، تعداد فایل، تعداد دانش‌آموز، تعداد فایل ناموجود)."""
    n_files = n_students = missing = 0
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for s in repo.students_for(aid):
            last = repo.latest_final(aid, s["id"])
            if not last:
                continue
            n_students += 1
            for i, f in enumerate(repo.sub_files(last["id"]), 1):
                lp = Path(f["local_path"]) if f["local_path"] else None
                if not lp or not lp.exists():
                    lp = cfg.files_dir / f"a{aid}" / f"re_{last['id']}_{i}_{_safe(f['name'] or 'file')}"
                    if not await bale.download(f["file_id"], lp):
                        missing += 1
                        continue
                    repo.x("UPDATE submission_files SET local_path=? WHERE id=?", (str(lp), f["id"]))
                z.write(lp, f"{_safe(s['full_name'])}/{i}_{_safe(f['name'] or lp.name)}")
                n_files += 1
    return path, n_files, n_students, missing
