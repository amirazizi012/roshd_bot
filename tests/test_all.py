import asyncio
import tempfile
import unittest
from datetime import datetime, timedelta, timezone
from pathlib import Path
from zoneinfo import ZoneInfo

from roshd_bot import rules
from roshd_bot.bot import App
from roshd_bot.config import Config
from roshd_bot.db import Repo
from roshd_bot.timeutil import en, fa, fmt_duration, parse_deadline, utcnow

TZ = ZoneInfo("Asia/Tehran")
T0 = datetime(2026, 10, 7, 12, 0, tzinfo=timezone.utc)


def rec(i, status, kind="regular", early=False, pts=0):
    dl = T0 + timedelta(days=i)
    sub = dl - timedelta(hours=13 if early else 1) if status in (rules.ON_TIME, rules.LATE) else None
    return rules.Record(kind, dl, status, sub, pts)


class RulesTest(unittest.TestCase):
    def test_classify(self):
        dl = T0
        c = lambda s, **k: rules.classify(s, dl, 20, T0 + timedelta(days=1), **k)
        self.assertEqual(c(dl - timedelta(minutes=1)).status, rules.ON_TIME)
        o = c(dl + timedelta(minutes=20))
        self.assertEqual((o.status, o.in_grace), (rules.ON_TIME, True))
        o = c(dl + timedelta(minutes=21))
        self.assertEqual((o.status, o.late_seconds), (rules.LATE, 21 * 60))
        self.assertEqual(c(dl + timedelta(hours=3), excused=True).status, rules.EXCUSED)
        self.assertEqual(c(None).status, rules.MISSING)
        self.assertEqual(rules.classify(None, dl, 20, dl + timedelta(minutes=10)).status, rules.OPEN)

    def test_streak_five_with_points(self):
        s = rules.summarize([rec(i, rules.ON_TIME) for i in range(4)])
        self.assertNotIn("commit:1", s.badges)
        s = rules.summarize([rec(i, rules.ON_TIME) for i in range(5)])
        self.assertIn("commit:1", s.badges)
        self.assertEqual(s.points, rules.STREAK_POINTS)
        s = rules.summarize([rec(i, rules.ON_TIME) for i in range(10)])
        self.assertEqual(s.points, 2 * rules.STREAK_POINTS)

    def test_one_late_breaks(self):
        rs = [rec(i, rules.ON_TIME) for i in range(4)] + [rec(4, rules.LATE)] + [rec(i, rules.ON_TIME) for i in range(5, 9)]
        s = rules.summarize(rs)
        self.assertNotIn("commit:1", s.badges)
        self.assertEqual(s.run, 4)

    def test_twelve_two_badges_and_missing_breaks(self):
        s = rules.summarize([rec(i, rules.ON_TIME) for i in range(10)])
        self.assertTrue({"commit:1", "commit:2"} <= set(s.badges))
        s = rules.summarize([rec(i, rules.ON_TIME) for i in range(3)] + [rec(3, rules.MISSING)] + [rec(i, rules.ON_TIME) for i in range(4, 8)])
        self.assertNotIn("commit:1", s.badges)

    def test_excused_and_optional_are_neutral(self):
        rs = [rec(i, rules.ON_TIME) for i in range(3)] + [rec(3, rules.EXCUSED), rec(4, rules.LATE, kind="optional")] \
             + [rec(i, rules.ON_TIME) for i in range(5, 7)]
        self.assertIn("commit:1", rules.summarize(rs).badges)

    def test_other_badges(self):
        s = rules.summarize([rec(i, rules.ON_TIME, early=True) for i in range(5)])
        self.assertIn("early", s.badges)
        self.assertIn("first_step", s.badges)
        s = rules.summarize([rec(0, rules.LATE)] + [rec(i, rules.ON_TIME) for i in range(1, 4)])
        self.assertIn("comeback", s.badges)
        s = rules.summarize([rec(i, rules.ON_TIME, kind="bonus", pts=5) for i in range(3)])
        self.assertEqual((s.points, "explorer" in s.badges), (15, True))


class TimeTest(unittest.TestCase):
    def test_digits_duration(self):
        self.assertEqual(fa("12:30"), "۱۲:۳۰")
        self.assertEqual(en("۱۴۰۵/۰۸/۱۵"), "1405/08/15")
        self.assertEqual(fmt_duration(2 * 3600 + 15 * 60), "۲ ساعت و ۱۵ دقیقه")
        self.assertEqual(fmt_duration(30), "کمتر از یک دقیقه")

    def test_parse_deadline(self):
        d = parse_deadline("۱۴۰۵/۰۸/۲۰ ۲۱:۰۰", TZ, T0)
        self.assertEqual(d.astimezone(TZ).strftime("%Y-%m-%d %H:%M"), "2026-11-11 21:00")
        d = parse_deadline("فردا ۲۱:۰۰", TZ, T0)
        self.assertEqual(d.astimezone(TZ).strftime("%Y-%m-%d %H:%M"), "2026-10-08 21:00")
        d = parse_deadline("پس‌فردا 9:30", TZ, T0)
        self.assertEqual(d.astimezone(TZ).day, 9)
        for bad in ("سلام", "۱۴۰۵/۱۳/۴۰ ۲۱:۰۰", "فردا ۲۵:۰۰"):
            with self.assertRaises(ValueError):
                parse_deadline(bad, TZ, T0)


class FakeBale:
    def __init__(self):
        self.sent = []   # (chat, text, markup)
        self.uid = 0

    async def send_message(self, chat, text, kb=None):
        self.sent.append((chat, text, kb)); return {}

    async def edit_message_text(self, chat, mid, text, kb=None):
        self.sent.append((chat, text, kb)); return {}

    async def answer_callback(self, *a, **k): pass
    async def send_photo(self, chat, fid, caption=None): self.sent.append((chat, f"[photo {fid}]", None))
    async def send_document(self, chat, fid, caption=None): self.sent.append((chat, f"[doc {fid}]", None))
    async def send_file(self, chat, path, kind="document", caption=None):
        self.sent.append((chat, f"[upload {kind} {Path(path).name}] {caption or ''}", None))
        return {"photo": [{"file_id": "LOGO"}]}

    async def download(self, file_id, dest):
        dest.parent.mkdir(parents=True, exist_ok=True); dest.write_bytes(b"x"); return True

    def last(self, chat):
        return [s for s in self.sent if s[0] == chat][-1]

    def buttons(self, chat):
        kbs = [m[2] for m in self.sent if m[0] == chat and m[2] and "inline_keyboard" in m[2]]
        kb = kbs[-1] if kbs else {}
        return [(b["text"], b["callback_data"]) for row in kb.get("inline_keyboard", []) for b in row]


class E2E(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        tmp = Path(tempfile.mkdtemp())
        self.cfg = Config("t", frozenset({1}), tmp, "x", 20, TZ, "رشد")
        self.bale, self.repo = FakeBale(), Repo(":memory:")
        self.app = App(self.cfg, self.bale, self.repo)
        self.n = 0

    async def msg(self, chat, text=None, **extra):
        self.n += 1
        m = {"chat": {"id": chat, "type": "private"}, "message_id": self.n, "date": int(utcnow().timestamp()), **extra}
        if text is not None:
            m["text"] = text
        await self.app.handle_update({"update_id": self.n, "message": m})

    async def press(self, chat, label):
        data = next(d for t, d in self.bale.buttons(chat) if label in t)
        self.n += 1
        await self.app.handle_update({"update_id": self.n, "callback_query": {
            "id": str(self.n), "data": data, "message": {"chat": {"id": chat}, "message_id": 1}}})

    async def test_full_flow(self):
        A, S = 1, 500
        await self.msg(A, "/start")
        # گروه و دانش‌آموز
        await self.msg(A, "🗂 گروه‌ها"); await self.press(A, "گروه جدید"); await self.msg(A, "ششم الف")
        await self.msg(A, "👥 دانش‌آموزان"); await self.press(A, "افزودن یک نفر")
        await self.msg(A, "علی رضایی"); await self.press(A, "ششم الف"); await self.press(A, "ادامه")
        code = self.repo.students()[0]["code"]
        self.assertIn(fa(code), self.bale.last(A)[1])
        # ثبت‌نام دانش‌آموز
        await self.msg(S, "/start"); await self.press(S, "ششم الف"); await self.press(S, "علی")
        await self.msg(S, "۰۰۰۰۰۰")
        self.assertIn("درست نیست", self.bale.last(S)[1])
        await self.msg(S, fa(code))
        self.assertIn("ثبت‌نام شد", self.bale.last(S)[1])
        # ساخت تکلیف
        await self.msg(A, "➕ تکلیف جدید"); await self.msg(A, "فصل ۱ ریاضی"); await self.msg(A, "تمرین صفحه ۱۰")
        await self.msg(A, photo=[{"file_id": "P1"}]); await self.press(A, "ادامه"); await self.press(A, "تکلیف اصلی")
        await self.press(A, "همهٔ دانش‌آموزان"); await self.press(A, "ادامه")
        await self.msg(A, "فردا ۲۱:۰۰"); await self.press(A, "انتشار")
        self.assertIn("برای ۱ نفر", self.bale.last(A)[1])
        self.assertIn("تکلیف جدید", self.bale.last(S)[1])
        # دانش‌آموز می‌بیند و ارسال می‌کند
        await self.press(S, "مشاهده")
        self.assertTrue(any("[photo P1]" in s[1] for s in self.bale.sent if s[0] == S))
        await self.press(S, "ارسال پاسخ")
        await self.msg(S, photo=[{"file_id": "A1"}]); await self.msg(S, photo=[{"file_id": "A2"}])
        await self.press(S, "ثبت نهایی")
        self.assertIn("به‌موقع", self.bale.last(S)[1])
        self.assertIn("اولین گام", self.bale.last(S)[1])
        self.assertEqual(len(self.repo.sub_files(1)), 2)
        # گزارش مدیر
        await self.msg(A, "📋 تکالیف"); await self.press(A, "فصل")
        self.assertIn("علی رضایی", self.bale.last(A)[1])
        self.assertIn("✅ به‌موقع (۱)", self.bale.last(A)[1])

    async def test_late_submission_recorded(self):
        gid = self.repo.add_group("g")
        sid, code = self.repo.add_student("سارا", [gid])
        self.repo.link_student(sid, 600)
        dl = utcnow() + timedelta(hours=1)
        aid = self.repo.create_assignment("t", "-", "regular", dl.strftime("%Y-%m-%dT%H:%M:%S"), 20, 0, [("all", 0)], [])
        a = self.repo.assignment(aid)
        sub = self.repo.get_or_create_draft(aid, sid)
        late_t = dl + timedelta(hours=2, minutes=15)
        self.repo.add_sub_file(sub, "photo", "f", None, "p", late_t)
        self.repo.finalize(sub)
        o, _, _ = self.repo.evaluate(a, sid, late_t)
        self.assertEqual((o.status, o.late_seconds), (rules.LATE, 2 * 3600 + 15 * 60))
        # تمدید مهلت برای همین دانش‌آموز، وضعیت را به‌موقع می‌کند
        ext = (dl + timedelta(hours=5)).strftime("%Y-%m-%dT%H:%M:%S")
        self.repo.x("INSERT INTO deadline_overrides VALUES(?,?,?)", (aid, sid, ext))
        self.assertEqual(self.repo.evaluate(a, sid, late_t)[0].status, rules.ON_TIME)

    async def test_late_joiner_not_penalized(self):
        gid = self.repo.add_group("g")
        old = utcnow() - timedelta(days=3)
        self.repo.create_assignment("old", "-", "regular", (old).strftime("%Y-%m-%dT%H:%M:%S"), 20, 0, [("all", 0)], [])
        self.repo.x("UPDATE assignments SET created_at=?", ((old - timedelta(days=1)).strftime("%Y-%m-%dT%H:%M:%S"),))
        sid, _ = self.repo.add_student("نو", [gid])
        self.assertEqual(self.repo.records(sid, utcnow()), [])


class Phase3(E2E):
    async def seed(self):
        gid = self.repo.add_group("g")
        sid, _ = self.repo.add_student("سارا", [gid]); self.repo.link_student(sid, 600)
        dl = utcnow() - timedelta(hours=2)
        aid = self.repo.create_assignment("هندسه", "-", "regular", dl.strftime("%Y-%m-%dT%H:%M:%S"), 20, 0, [("all", 0)], [])
        return sid, aid, dl

    def sub(self, aid, sid, t, local=None):
        d = self.repo.get_or_create_draft(aid, sid)
        self.repo.add_sub_file(d, "photo", "f", local, "p.jpg", t)
        self.repo.finalize(d)

    async def test_last_send_time_counts(self):
        sid, aid, dl = await self.seed()
        a = self.repo.assignment(aid)
        self.sub(aid, sid, dl - timedelta(hours=1))
        self.sub(aid, sid, dl + timedelta(minutes=10))          # جایگزینی در مهلت ارفاقی → آخرین زمان حساب می‌شود
        o, at, _ = self.repo.evaluate(a, sid, utcnow())
        self.assertEqual((o.status, o.in_grace, at), (rules.ON_TIME, True, dl + timedelta(minutes=10)))

    async def test_resubmit_via_admin_not_penalized_but_first_late_is_late(self):
        sid, aid, dl = await self.seed()
        a = self.repo.assignment(aid)
        self.sub(aid, sid, dl - timedelta(hours=1))
        self.repo.request_resubmit(aid, sid, "تار")
        self.sub(aid, sid, utcnow())                            # بعد از مهلت، داخل پنجرهٔ مدیر
        self.assertEqual(self.repo.evaluate(a, sid, utcnow())[0].status, rules.ON_TIME)
        # ارسال اول دیرهنگام: دیر می‌ماند
        sid2, _ = self.repo.add_student("رضا", [1]); self.repo.link_student(sid2, 601)
        self.sub(aid, sid2, dl + timedelta(hours=1))
        self.assertEqual(self.repo.evaluate(a, sid2, utcnow())[0].status, rules.LATE)

    async def test_panel(self):
        from aiohttp.test_utils import TestClient, TestServer
        from dataclasses import replace
        from roshd_bot.panel import make_app
        sid, aid, dl = await self.seed()
        fp = self.cfg.files_dir / "a1" / "x.jpg"; fp.parent.mkdir(parents=True, exist_ok=True); fp.write_bytes(b"IMG")
        self.sub(aid, sid, dl - timedelta(minutes=5), str(fp))
        outside = self.cfg.data_dir / "secret.txt"; outside.write_text("s")
        d = self.repo.get_or_create_draft(aid, sid)
        self.repo.add_sub_file(d, "document", "g", str(outside), "s.txt", utcnow())
        cfg = replace(self.cfg, panel_password="pw-123456")
        async with TestClient(TestServer(make_app(cfg, self.repo))) as c:
            r = await c.get("/", allow_redirects=False); self.assertEqual((r.status, r.headers["Location"]), (302, "/login"))
            r = await c.post("/login", data={"password": "bad"}); self.assertIn("رمز درست نیست", await r.text())
            r = await c.post("/login", data={"password": "pw-123456"}, allow_redirects=False); self.assertEqual(r.status, 302)
            r = await c.get("/"); t = await r.text(); self.assertEqual(r.status, 200); self.assertIn("سارا", t)
            r = await c.get(f"/assignment/{aid}"); self.assertIn("هندسه", await r.text())
            r = await c.get(f"/student/{sid}"); self.assertEqual(r.status, 200)
            r = await c.get("/file/1"); self.assertEqual((r.status, await r.read()), (200, b"IMG"))
            r = await c.get("/file/2"); self.assertEqual(r.status, 404)          # خارج از پوشهٔ فایل‌ها
            r = await c.get("/export.xlsx"); self.assertEqual(r.status, 200); self.assertTrue((await r.read()).startswith(b"PK"))
            for path in ("/", f"/assignment/{aid}", "/file/1"):                  # فقط‌خواندنی
                r = await c.post(path, data={}); self.assertEqual(r.status, 405)
        async with TestClient(TestServer(make_app(cfg, self.repo))) as c:      # قفل بعد از ۵ تلاش اشتباه
            for _ in range(5):
                await c.post("/login", data={"password": "x"})
            r = await c.post("/login", data={"password": "pw-123456"}); self.assertIn("تلاش‌های ناموفق", await r.text())


class Phase2(E2E):
    async def setup_one(self, deadline_delta=timedelta(hours=1)):
        gid = self.repo.add_group("g")
        sid, code = self.repo.add_student("سارا", [gid])
        self.repo.link_student(sid, 600 if sid == 1 else 600 + sid)
        dl = utcnow() + deadline_delta
        aid = self.repo.create_assignment("هندسه", "-", "regular", dl.strftime("%Y-%m-%dT%H:%M:%S"), 20, 0, [("all", 0)], [])
        return gid, sid, aid, dl

    def submit_at(self, aid, sid, t, fid="f"):
        sub = self.repo.get_or_create_draft(aid, sid)
        self.repo.add_sub_file(sub, "photo", fid, None, "p.jpg", t)
        self.repo.finalize(sub)

    async def test_hard_close_24h(self):
        _, sid, aid, dl = await self.setup_one(timedelta(hours=-25))   # مهلت ۲۵ ساعت پیش
        a = self.repo.assignment(aid)
        self.assertFalse(self.repo.can_submit(a, sid, utcnow()))
        _, sid2, aid2, dl2 = await self.setup_one(timedelta(hours=-23))
        self.assertTrue(self.repo.can_submit(self.repo.assignment(aid2), sid2, utcnow()))

    async def test_replace_then_resubmit_window(self):
        _, sid, aid, dl = await self.setup_one(timedelta(hours=-2))
        a = self.repo.assignment(aid)
        self.submit_at(aid, sid, dl - timedelta(minutes=5))
        self.assertFalse(self.repo.can_submit(a, sid, utcnow()))            # بعد از مهلت، بدون درخواست: بسته
        self.repo.request_resubmit(aid, sid, "تار")
        self.assertTrue(self.repo.can_submit(a, sid, utcnow()))             # با درخواست مدیر: باز
        self.assertTrue(self.repo.can_submit(a, sid, utcnow() + timedelta(hours=23)))
        self.assertFalse(self.repo.can_submit(a, sid, utcnow() + timedelta(hours=25)))
        # ارسال مجدد وضعیت زمانی را عوض نمی‌کند
        self.submit_at(aid, sid, utcnow(), "f2")
        self.assertEqual(self.repo.evaluate(a, sid, utcnow())[0].status, rules.ON_TIME)
        self.assertEqual(len(self.repo.final_attempts(aid, sid)), 2)

    async def test_excuse_and_late_neutral(self):
        _, sid, aid, dl = await self.setup_one(timedelta(hours=-3))
        a = self.repo.assignment(aid)
        self.assertEqual(self.repo.evaluate(a, sid, utcnow())[0].status, rules.MISSING)
        self.assertTrue(self.repo.toggle_excuse(aid, sid))
        self.assertEqual(self.repo.evaluate(a, sid, utcnow())[0].status, rules.EXCUSED)
        self.assertFalse(self.repo.toggle_excuse(aid, sid))

    async def test_admin_panel_flow(self):
        A = 1
        _, sid, aid, dl = await self.setup_one()
        self.submit_at(aid, sid, utcnow(), "ANS1")
        await self.msg(A, "📋 تکالیف"); await self.press(A, "هندسه")
        await self.press(A, "بررسی دانش‌آموز"); await self.press(A, "سارا")
        await self.press(A, "دیدن پاسخ")
        self.assertTrue(any("[photo ANS1]" in x[1] for x in self.bale.sent if x[0] == A))
        await self.press(A, "بازخورد"); await self.msg(A, "عالی بود")
        self.assertIn("عالی بود", self.bale.last(600)[1])
        await self.press(A, "ارسال مجدد"); await self.msg(A, "-")
        self.assertIn("دوباره بفرستی", self.bale.last(600)[1])
        await self.press(A, "تمدید فردی"); await self.msg(A, "پس‌فردا ۲۱:۰۰")
        self.assertIn("تمدید شد", self.bale.last(600)[1])
        await self.press(A, "عذر موجه")
        self.assertTrue(self.repo.is_excused(aid, sid))
        # دانش‌آموز بازخورد را می‌بیند
        await self.press(600, "مشاهده")
        self.assertTrue(any("بازخورد مدیر: عالی بود" in x[1] for x in self.bale.sent if x[0] == 600))

    async def test_broadcast_edit_extend_archive(self):
        A = 1
        gid, sid, aid, dl = await self.setup_one()
        await self.msg(A, "📢 پیام همگانی"); await self.press(A, "همهٔ دانش‌آموزان"); await self.press(A, "ادامه")
        await self.msg(A, "کلاس فردا تعطیل است")
        self.assertIn("کلاس فردا", self.bale.last(600)[1])
        await self.msg(A, "📋 تکالیف"); await self.press(A, "هندسه")
        await self.press(A, "ویرایش"); await self.press(A, "عنوان"); await self.msg(A, "هندسهٔ ۲")
        self.assertEqual(self.repo.assignment(aid)["title"], "هندسهٔ ۲")
        await self.msg(A, "📋 تکالیف"); await self.press(A, "هندسه")
        await self.press(A, "تمدید برای همه"); await self.msg(A, "پس‌فردا ۲۲:۰۰")
        self.assertIn("تمدید شد", self.bale.last(600)[1])
        await self.press(A, "بایگانی")
        self.assertEqual(self.repo.assignment(aid)["status"], "archived")
        self.assertEqual(self.repo.assignments_for(sid), [])

    async def test_logo_cached(self):
        gid = self.repo.add_group("g"); self.repo.add_student("x", [gid])
        await self.msg(900, "/start"); await self.msg(901, "/start")
        self.assertEqual(self.repo.meta_get("logo_file_id"), "LOGO")
        uploads = [x for x in self.bale.sent if "[upload photo" in x[1]]
        self.assertEqual(len(uploads), 1)

    async def test_excel_and_zip(self):
        from openpyxl import load_workbook
        from roshd_bot.export import build_excel, build_zip
        gid, sid, aid, dl = await self.setup_one(timedelta(hours=-1))
        sub = self.repo.get_or_create_draft(aid, sid)
        fp = self.cfg.data_dir / "ans.jpg"; fp.parent.mkdir(parents=True, exist_ok=True); fp.write_bytes(b"x")
        self.repo.add_sub_file(sub, "photo", "f", str(fp), "p.jpg", dl + timedelta(hours=1)); self.repo.finalize(sub)
        x = build_excel(self.repo, self.cfg, self.cfg.data_dir / "s.xlsx", utcnow())
        wb = load_workbook(x)
        self.assertEqual(wb.sheetnames, ["خلاصه دانش‌آموزان", "ماتریس تکالیف", "جزئیات ارسال"])
        det = list(wb["جزئیات ارسال"].iter_rows(min_row=5, values_only=True))[0]
        self.assertEqual((det[3], det[4]), ("سارا", "با تاخیر"))
        self.assertEqual(det[6], 60)
        import zipfile
        p, nf, ns, miss = await build_zip(self.repo, self.bale, self.cfg, aid, self.cfg.data_dir / "z.zip")
        self.assertEqual((nf, ns, miss), (1, 1, 0))
        self.assertTrue(any(n.startswith("سارا/") for n in zipfile.ZipFile(p).namelist()))
        await self.msg(1, "📊 خروجی اکسل")
        self.assertTrue(any("assignment_stats.xlsx" in y[1] for y in self.bale.sent))


if __name__ == "__main__":
    unittest.main()
